#!/usr/bin/env python3
"""Export the 2027 THEI grid 'Plan Info' tab (carrier vendors, HRA, contacts) to the Max KB.

Reads the same Drive .xlsx export as export_2027_grid_to_kb.py (/tmp/thei-2027-grid.xlsx)
and writes max-knowledge/carriers/plan-info-2027.md: one section per carrier, one line per
non-blank column (Dental, Vision, HRA Payment, OTC, Transportation, ...). Read-only on the sheet.
Re-run after each sheet refresh.

Dentist network checks still use data/dental-networks-2027.json (curated from this tab's
Dental column); this script prints a warning when a carrier's Dental cell is not reflected there.
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from datetime import datetime, timezone
from pathlib import Path

import openpyxl

REPO_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(Path(__file__).resolve().parent))
from thei_grid_common import SHEET_ID_2027 as SHEET_ID, require_xlsx  # noqa: E402

OUT_PATH = REPO_ROOT / "max-knowledge" / "carriers" / "plan-info-2027.md"
DENTAL_JSON = REPO_ROOT / "data" / "dental-networks-2027.json"


def clean(v) -> str:
    if v is None:
        return ""
    s = str(v).replace("\r", "")
    s = "\n".join(re.sub(r"[ \t]+", " ", ln).strip() for ln in s.split("\n"))
    s = re.sub(r"\n{2,}", "\n", s).strip()
    return s


def parse(xlsx: Path) -> tuple[list[str], list[dict]]:
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    if "Plan Info" not in wb.sheetnames:
        raise SystemExit("no 'Plan Info' tab in the workbook")
    ws = wb["Plan Info"]
    rows = list(ws.iter_rows(values_only=True))
    header = [clean(h).replace("\n", " ") for h in rows[0]]
    cols = [i for i, h in enumerate(header) if h and h != "■"]
    carriers = []
    for r in rows[1:]:
        name = clean(r[0]).replace("\n", " ") if r and r[0] else ""
        if not name:
            continue
        fields = {header[i]: clean(r[i]) for i in cols[1:] if i < len(r) and clean(r[i])}
        carriers.append({"carrier": name, "fields": fields})
    return [header[i] for i in cols], carriers


def render(columns: list[str], carriers: list[dict], pulled: str) -> str:
    out = [
        "# 2027 Plan Info by carrier (THEI 2027 Plan Comparison Grid, 'Plan Info' tab)",
        "",
        f"Workbook: https://docs.google.com/spreadsheets/d/{SHEET_ID}/edit",
        f"Pulled: {pulled}",
        "",
        "Carrier-level reference: broker managers, broker support, enrollment portals, HRA payment, "
        "member services, and the vendors for fitness, lab, pharmacy, mail order, M3P, vision, dental, "
        "diabetic supplies, OTC, transportation, hearing, DME, healthy food and utilities. "
        "For 2027 questions this tab is newer than `carrier-contacts.md` (2026); when they disagree, cite this file. "
        "A blank column means the grid has nothing for that carrier: say it is not on the grid, do not guess. "
        "Plan-level dollars (OTC amounts, dental allowances, copays) are in PLAN DATA and the `*-plans-florida-2027` docs, not here.",
        "",
    ]
    for c in carriers:
        out += [f"## {c['carrier']}", ""]
        for col in columns[1:]:
            v = c["fields"].get(col)
            if not v:
                continue
            lines = v.split("\n")
            if len(lines) == 1:
                out.append(f"- **{col}:** {lines[0]}")
            else:
                out.append(f"- **{col}:**")
                out += [f"  - {ln}" for ln in lines if ln]
        out.append("")
    return "\n".join(out).rstrip() + "\n"


def check_dental(carriers: list[dict]) -> list[str]:
    warn = []
    if not DENTAL_JSON.exists():
        return ["data/dental-networks-2027.json missing"]
    d = json.loads(DENTAL_JSON.read_text())
    blob = json.dumps(d).lower()
    for c in carriers:
        dental = c["fields"].get("Dental", "")
        if not dental:
            warn.append(f"{c['carrier']}: Dental cell blank on Plan Info")
            continue
        key = re.split(r"[\n(]", dental)[0].strip().lower()
        key = re.sub(r"\s+", " ", key)
        probe = {"denta quest": "dentaquest", "uhc": "uhc"}.get(key, key)
        if probe and probe.split(" ")[0] not in blob:
            warn.append(f"{c['carrier']}: Dental '{key}' not found in dental-networks-2027.json")
    return warn


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--xlsx", type=Path, default=Path("/tmp/thei-2027-grid.xlsx"))
    args = ap.parse_args()
    require_xlsx(args.xlsx)
    columns, carriers = parse(args.xlsx)
    pulled = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    OUT_PATH.write_text(render(columns, carriers, pulled), encoding="utf-8")
    print(f"wrote {OUT_PATH} ({len(carriers)} carriers, {len(columns) - 1} columns)")
    for w in check_dental(carriers):
        print("WARN", w)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
